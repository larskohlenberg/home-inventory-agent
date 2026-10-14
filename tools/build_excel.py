#!/usr/bin/env python3
"""Baut Inventar.xlsx aus data/raeume/*.json.

Korrekturspalten (Korr. …, Löschen, Geprüft, Bemerkung) aus einer vorhandenen
Inventar.xlsx werden per ID übernommen und zusätzlich in data/korrekturen.json
gesichert – Nutzereinträge gehen beim Neuaufbau nie verloren.

Aufruf: python tools/build_excel.py <Projektordner>   (Standard: aktueller Ordner)
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PROJ = Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve()
XLSX = PROJ / "Inventar.xlsx"
RAEUME = PROJ / "data" / "raeume"
KORR_JSON = PROJ / "data" / "korrekturen.json"

# (Überschrift, Feld/Typ, Breite)
COLS = [
    ("ID", "id", 12),
    ("Gegenstand", "gegenstand", 34),
    ("Marke/Modell", "marke_modell", 30),
    ("Kategorie", "kategorie", 20),
    ("Zuordnung", "zuordnung", 11),
    ("Anzahl", "anzahl", 8),
    ("Einzelpreis €", "einzelpreis", 13),
    ("Gesamt €", "=GESAMT", 13),
    ("Sicherheit", "sicherheit", 10),
    ("Fundort", "fundort", 24),
    ("Hinweis", "hinweis", 30),
    ("Korr. Marke/Modell", "K:marke_modell", 22),
    ("Korr. Preis €", "K:einzelpreis", 12),
    ("Korr. Anzahl", "K:anzahl", 10),
    ("Korr. Zuordnung", "K:zuordnung", 12),
    ("Löschen (x)", "K:loeschen", 9),
    ("Geprüft (x)", "K:geprueft", 9),
    ("Bemerkung", "K:bemerkung", 28),
    ("Preisquelle", "=LINK_PREIS", 18),
    ("Preis-Auszug", "preis_auszug", 40),
    ("Abrufdatum", "abrufdatum", 11),
    ("Nachweis", "=LINK_NACHWEIS", 11),
    ("Standbild", "=LINK_STANDBILD", 11),
    ("Video-Zeit", "video_zeit", 9),
    ("Gesehen bei", "=GESEHEN", 20),
    ("Sammelposten", "sammelposten", 11),
    ("Zuordnung wirksam", "=ZUORDNUNG", 12),
]
COL = {h: get_column_letter(i + 1) for i, (h, _, _) in enumerate(COLS)}
KORR_FIELDS = [f[2:] for _, f, _ in COLS if f.startswith("K:")]

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
KORR_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FONT = Font(bold=True, color="FFFFFF")
EUR = '#,##0.00 "€"'
LINK = Font(color="0563C1", underline="single")


def sheet_title(raum: str) -> str:
    return "".join(c for c in raum if c not in r"[]:*?/\\")[:31]


def load_existing_corrections() -> dict:
    korr = json.loads(KORR_JSON.read_text(encoding="utf-8")) if KORR_JSON.exists() else {}
    if XLSX.exists():
        wb = load_workbook(XLSX)
        for ws in wb.worksheets:
            if ws.title == "Übersicht":
                continue
            header = [c.value for c in ws[1]]
            if "ID" not in header:
                continue
            idx = {h: i for i, h in enumerate(header)}
            for row in ws.iter_rows(min_row=2, values_only=True):
                pid = row[idx["ID"]]
                if not pid:
                    continue
                vals = {}
                for h, f, _ in COLS:
                    if f.startswith("K:") and h in idx and row[idx[h]] not in (None, ""):
                        vals[f[2:]] = row[idx[h]]
                if vals:
                    korr[pid] = vals
                else:
                    korr.pop(pid, None)
    return korr


def write_room(ws, room: dict, korr: dict):
    ws.append([h for h, _, _ in COLS])
    for i, (h, f, w) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"

    for r, p in enumerate(room["positionen"], start=2):
        k = korr.get(p["id"], {})
        for c, (h, f, _) in enumerate(COLS, 1):
            cell = ws.cell(row=r, column=c)
            if f == "=GESAMT":
                cell.value = (f'=IF({COL["Löschen (x)"]}{r}<>"",0,'
                              f'IF({COL["Korr. Preis €"]}{r}="",{COL["Einzelpreis €"]}{r},{COL["Korr. Preis €"]}{r})*'
                              f'IF({COL["Korr. Anzahl"]}{r}="",{COL["Anzahl"]}{r},{COL["Korr. Anzahl"]}{r}))')
                cell.number_format = EUR
                cell.font = Font(bold=True)
            elif f == "=ZUORDNUNG":
                cell.value = f'=IF({COL["Korr. Zuordnung"]}{r}="",{COL["Zuordnung"]}{r},{COL["Korr. Zuordnung"]}{r})'
            elif f == "=LINK_PREIS":
                if p.get("preis_url"):
                    cell.value, cell.hyperlink, cell.font = p.get("preis_quelle") or "Link", p["preis_url"], LINK
            elif f == "=LINK_NACHWEIS":
                if p.get("nachweis") and (PROJ / p["nachweis"]).exists():
                    cell.value, cell.hyperlink, cell.font = "Screenshot", p["nachweis"], LINK
            elif f == "=LINK_STANDBILD":
                if p.get("standbild") and (PROJ / p["standbild"]).exists():
                    cell.value, cell.hyperlink, cell.font = "Bild", p["standbild"], LINK
            elif f.startswith("K:"):
                cell.value = k.get(f[2:])
                cell.fill = KORR_FILL
            elif f == "=GESEHEN":
                cell.value = ", ".join(p.get("gesehen_bei") or [])
            elif f == "sammelposten":
                cell.value = "ja" if p.get("sammelposten") else ""
            else:
                cell.value = p.get(f)
            if h in ("Einzelpreis €", "Korr. Preis €"):
                cell.number_format = EUR
    last = len(room["positionen"]) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{last}"
    ws.column_dimensions[COL["Zuordnung wirksam"]].hidden = True

    dv_x = DataValidation(type="list", formula1='"x"', allow_blank=True)
    dv_z = DataValidation(type="list", formula1='"Hausrat,Gebäude,Unklar"', allow_blank=True)
    ws.add_data_validation(dv_x)
    ws.add_data_validation(dv_z)
    for h in ("Löschen (x)", "Geprüft (x)"):
        dv_x.add(f"{COL[h]}2:{COL[h]}{max(last, 2) + 200}")
    dv_z.add(f"{COL['Korr. Zuordnung']}2:{COL['Korr. Zuordnung']}{max(last, 2) + 200}")


def write_overview(ws, rooms: list[dict]):
    ws["A1"] = "Hausratinventar – Übersicht"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Stand: {date.today():%d.%m.%Y} · Werte = Neuwert (Wiederbeschaffung) · gelbe Spalten in den Raumblättern = deine Korrekturen"
    ws["A2"].font = Font(italic=True, color="666666")

    heads = ["Raum", "Hausrat €", "Gebäude €", "Unklar €", "Positionen", "geprüft"]
    ws.append([])
    ws.append(heads)
    hr = ws.max_row
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=hr, column=i)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    G, Z, P = COL["Gesamt €"], COL["Zuordnung wirksam"], COL["Geprüft (x)"]
    for room in rooms:
        t = f"'{sheet_title(room['raum'])}'"
        ws.append([
            room["raum"],
            f'=SUMIFS({t}!{G}:{G},{t}!{Z}:{Z},"Hausrat")',
            f'=SUMIFS({t}!{G}:{G},{t}!{Z}:{Z},"Gebäude")',
            f'=SUMIFS({t}!{G}:{G},{t}!{Z}:{Z},"Unklar")',
            f'=COUNTA({t}!A:A)-1',
            f'=COUNTIF({t}!{P}:{P},"x")',
        ])
        ws.cell(row=ws.max_row, column=1).hyperlink = f"#{t}!A1"
        ws.cell(row=ws.max_row, column=1).font = LINK
    first, last = hr + 1, ws.max_row
    ws.append(["Summe"] + [f"=SUM({get_column_letter(c)}{first}:{get_column_letter(c)}{last})" for c in range(2, 7)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=first, max_row=ws.max_row, min_col=2, max_col=4):
        for c in row:
            c.number_format = EUR

    # Hausrat nach Kategorie über alle Räume
    ws.append([])
    ws.append(["Hausrat nach Kategorie", "Summe €"])
    kr = ws.max_row
    for i in (1, 2):
        ws.cell(row=kr, column=i).fill, ws.cell(row=kr, column=i).font = HEAD_FILL, HEAD_FONT
    K = COL["Kategorie"]
    cats = sorted({p["kategorie"] for r in rooms for p in r["positionen"]})
    for cat in cats:
        parts = [f"SUMIFS('{sheet_title(r['raum'])}'!{G}:{G},'{sheet_title(r['raum'])}'!{K}:{K},\"{cat}\","
                 f"'{sheet_title(r['raum'])}'!{Z}:{Z},\"Hausrat\")" for r in rooms]
        ws.append([cat, "=" + "+".join(parts) if parts else 0])
        ws.cell(row=ws.max_row, column=2).number_format = EUR

    for col, w in zip("ABCDEF", (34, 14, 14, 14, 11, 9)):
        ws.column_dimensions[col].width = w


def main():
    if not RAEUME.is_dir():
        sys.exit(f"Keine Raumdateien gefunden: {RAEUME} fehlt. Projektordner als Argument angeben.")
    rooms = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(RAEUME.glob("*.json"))]
    korr = load_existing_corrections()
    KORR_JSON.parent.mkdir(parents=True, exist_ok=True)
    KORR_JSON.write_text(json.dumps(korr, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    wb = Workbook()
    write_overview(wb.active, rooms)
    wb.active.title = "Übersicht"
    for room in rooms:
        write_room(wb.create_sheet(sheet_title(room["raum"])), room, korr)
    try:
        wb.save(XLSX)
    except PermissionError:
        sys.exit(f"{XLSX.name} ist noch geöffnet. Bitte schließen und erneut ausführen.")
    n = sum(len(r["positionen"]) for r in rooms)
    print(f"{XLSX.name}: {len(rooms)} Räume, {n} Positionen, {len(korr)} Korrekturen übernommen")


if __name__ == "__main__":
    main()
